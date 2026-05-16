'''
    文件做三类事:
        1. 文本清洗
        2. 文档切块与向量化入库
        3. 文档片段检索与上下文格式化
'''
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Literal
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree

from db.queries import (
    delete_document_chunks,             # 重建索引前先删旧 chunk
    find_similar_document_chunks,       # 片段检索
    save_document_chunk,                # 保存切块后的 chunk
    save_session_job_description,       # 保存 JD 原文记录
    save_session_resume,                # 保存简历原文记录
)
from rag.embeddings import embed_text
from rag.user_memories import index_resume_user_memories

DocumentSource = Literal["resume", "jd"]

DEFAULT_CHUNK_CHARS = 1000
DEFAULT_CHUNK_OVERLAP = 120
SUPPORTED_UPLOAD_EXTENSIONS = {".pdf", ".docx"}

# 文本清洗函数:清理换行、空格，保留简历/JD 的段落结构
def normalize_document_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    normalized = "\n".join(lines)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()

# 把长文本切成默认约 1000 字符的 chunks，超长段落带少量 overlap
def chunk_document_text(
    text: str,
    max_chars: int = DEFAULT_CHUNK_CHARS,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[str]:
    normalized = normalize_document_text(text)
    if not normalized:
        return []

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", normalized) if p.strip()]
    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(_split_long_text(paragraph, max_chars, overlap))
            continue

        candidate = f"{current}\n\n{paragraph}" if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                chunks.append(current)
            current = paragraph

    if current:
        chunks.append(current)

    return chunks

# 上传文件文本提取入口
def extract_uploaded_document_text(filename: str, file_bytes: bytes) -> str:
    """从上传的 PDF / DOCX 文件中提取纯文本；老式 .doc 暂不支持。"""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_text(file_bytes)
    if suffix == ".docx":
        return _extract_docx_text(file_bytes)
    if suffix == ".doc":
        raise ValueError("暂不支持旧版 .doc 文件，请转换为 .docx 或 PDF 后上传")
    raise ValueError("仅支持 PDF 或 DOCX 文件")

# 保存简历原文
# 生成 embedding
# 写入 document_chunks
# 成功后标记 processed，失败标记 failed
async def index_session_resume(
    session_id: str,
    user_id: str,
    content: str,
    source_type: str = "text",
    filename: str | None = None,
    metadata: dict | None = None,
    role: str | None = None,
) -> dict:
    normalized = normalize_document_text(content)
    if not normalized:
        raise ValueError("Resume content cannot be empty")
    # 在正式切块和向量化之前，先在数据库里保存一条简历记录，并把状态标成：status="pending"
    resume = save_session_resume(
        session_id=session_id,
        user_id=user_id,
        content=normalized,
        source_type=source_type,
        filename=filename,
        status="pending",
        metadata=metadata,
    )
    try:
        chunks = await _index_document_chunks(
            session_id=session_id,
            source="resume",
            content=normalized,
            resume_id=resume["id"],
        )
    except Exception:
        save_session_resume(
            session_id=session_id,
            user_id=user_id,
            content=normalized,
            source_type=source_type,
            filename=filename,
            status="failed",
            metadata=metadata,
        )
        raise

    resume = save_session_resume(
        session_id=session_id,
        user_id=user_id,
        content=normalized,
        source_type=source_type,
        filename=filename,
        status="processed",
        metadata=metadata,
    )
    # 简历上传成功后，系统还会顺手尝试从简历中抽取长期记忆
    try:
        await index_resume_user_memories(
            user_id=user_id,
            session_id=session_id,
            role=role,
            resume_content=normalized,
        )
    except Exception:
        pass
    return {"document": resume, "chunk_count": len(chunks)}

# 对象从简历换成 JD
async def index_session_job_description(
    session_id: str,
    user_id: str,
    content: str,
    source_type: str = "text",
    filename: str | None = None,
    metadata: dict | None = None,
) -> dict:
    normalized = normalize_document_text(content)
    if not normalized:
        raise ValueError("Job description content cannot be empty")

    job_description = save_session_job_description(
        session_id=session_id,
        user_id=user_id,
        content=normalized,
        source_type=source_type,
        filename=filename,
        status="pending",
        metadata=metadata,
    )
    try:
        chunks = await _index_document_chunks(
            session_id=session_id,
            source="jd",
            content=normalized,
            job_description_id=job_description["id"],
        )
    except Exception:
        save_session_job_description(
            session_id=session_id,
            user_id=user_id,
            content=normalized,
            source_type=source_type,
            filename=filename,
            status="failed",
            metadata=metadata,
        )
        raise

    job_description = save_session_job_description(
        session_id=session_id,
        user_id=user_id,
        content=normalized,
        source_type=source_type,
        filename=filename,
        status="processed",
        metadata=metadata,
    )
    return {"document": job_description, "chunk_count": len(chunks)}


async def retrieve_interview_document_context(
    session_id: str,
    query: str,
    threshold: float = 0.65,
    resume_limit: int = 4,
    jd_limit: int = 4,
    max_context_chars: int = 4000,
) -> dict:
    """检索和本次问题相关的简历/JD 片段，并格式化成面试官可直接使用的上下文。"""
    normalized_query = normalize_document_text(query)
    if not normalized_query:
        return {"resume_chunks": [], "job_description_chunks": [], "context": ""}

    query_embedding = await embed_text(normalized_query)
    resume_chunks = find_similar_document_chunks(
        session_id=session_id,
        query_embedding=query_embedding,
        threshold=threshold,
        limit=resume_limit,
        chunk_source="resume",
    )
    job_description_chunks = find_similar_document_chunks(
        session_id=session_id,
        query_embedding=query_embedding,
        threshold=threshold,
        limit=jd_limit,
        chunk_source="jd",
    )

    return {
        "resume_chunks": resume_chunks,
        "job_description_chunks": job_description_chunks,
        "context": format_document_context(
            resume_chunks=resume_chunks,
            job_description_chunks=job_description_chunks,
            max_chars=max_context_chars,
        ),
    }


def format_document_context(
    resume_chunks: list[dict],
    job_description_chunks: list[dict],
    max_chars: int = 4000,
) -> str:
    """把检索结果压成 prompt 友好的文本，避免把整份简历/JD 塞进上下文。"""
    sections: list[str] = []
    if resume_chunks:
        sections.append(_format_chunk_section("简历相关片段", resume_chunks))
    if job_description_chunks:
        sections.append(_format_chunk_section("岗位 JD 相关片段", job_description_chunks))

    context = "\n\n".join(sections).strip()
    if len(context) <= max_chars:
        return context
    return context[:max_chars].rstrip() + "\n..."

# 处理超长段落
def _split_long_text(text: str, max_chars: int, overlap: int) -> list[str]:
    chunks: list[str] = []
    start = 0
    safe_overlap = min(max(overlap, 0), max_chars // 2)

    while start < len(text):
        end = min(start + max_chars, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = end - safe_overlap

    return chunks

# PDF 提取函数
def _extract_pdf_text(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("缺少 pypdf 依赖，请先安装 backend/requirements.txt") from exc

    reader = PdfReader(BytesIO(file_bytes))
    page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(page_texts)
    normalized = normalize_document_text(text)
    if not normalized:
        raise ValueError("未能从 PDF 中提取到文本；如果是扫描件，请先转换为可复制文本")
    return normalized


def _extract_docx_text(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx 依赖，请先安装 backend/requirements.txt") from exc

    document = Document(BytesIO(file_bytes))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    normalized = normalize_document_text("\n".join(parts))
    if normalized:
        return normalized

    # 很多简历模板会把正文放在文本框、形状、页眉页脚等结构里，python-docx 可能读不到。
    # 这里直接扫描 DOCX 底层 Word XML 的 w:t 文本节点作为兜底。
    normalized = _extract_docx_xml_text(file_bytes)
    if not normalized:
        raise ValueError("未能从 DOCX 中提取到文本；如果内容在图片或扫描件中，请转为可复制文本后上传")
    return normalized

# DOCX 提取函数
def _extract_docx_xml_text(file_bytes: bytes) -> str:
    texts: list[str] = []
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            xml_names = [
                name for name in archive.namelist()
                if name.startswith("word/") and name.endswith(".xml")
            ]
            for name in xml_names:
                root = ElementTree.fromstring(archive.read(name))
                for element in root.iter():
                    # w:t 节点承载 Word 正文/文本框/页眉页脚等可见文本。
                    if element.tag.endswith("}t") and element.text:
                        texts.append(element.text)
    except (BadZipFile, ElementTree.ParseError) as exc:
        raise ValueError("DOCX 文件结构异常，请确认文件没有损坏") from exc

    return normalize_document_text("\n".join(texts))


def _format_chunk_section(title: str, chunks: list[dict]) -> str:
    lines = [f"{title}:"]
    for index, chunk in enumerate(chunks, start=1):
        similarity = chunk.get("similarity")
        score = f" similarity={similarity:.2f}" if isinstance(similarity, (int, float)) else ""
        lines.append(f"{index}. [{chunk.get('chunk_source', 'document')}{score}] {chunk.get('content', '')}")
    return "\n".join(lines)


# 先完成所有新 chunk 的 embedding
# 再删除旧 chunks 并写入新 chunks，避免中途状态不一致导致查询失败
async def _index_document_chunks(
    session_id: str,
    source: DocumentSource,
    content: str,
    resume_id: str | None = None,
    job_description_id: str | None = None,
) -> list[dict]:
    chunks = chunk_document_text(content)
    embedded_chunks: list[tuple[str, list[float]]] = []
    for chunk in chunks:
        embedded_chunks.append((chunk, await embed_text(chunk)))

    delete_document_chunks(session_id=session_id, chunk_source=source)

    saved_chunks: list[dict] = []
    for index, (chunk, embedding) in enumerate(embedded_chunks):
        saved_chunks.append(
            save_document_chunk(
                session_id=session_id,
                chunk_source=source,
                chunk_index=index,
                content=chunk,
                embedding=embedding,
                resume_id=resume_id,
                job_description_id=job_description_id,
                metadata={"chunk_chars": len(chunk)},
            )
        )

    return saved_chunks
